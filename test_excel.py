# -*- coding: utf-8 -*-
"""
테스트 데이터 3명분을 넣어서 엑셀이 제대로 계산되는지 검증
"""
import openpyxl
from openpyxl.utils import get_column_letter
import math

# 엑셀 열기
wb = openpyxl.load_workbook('충남_AHP_분석_결과_v3.xlsx')
ws = wb['응답데이터']

# 테스트 응답자 3명의 쌍대비교 데이터
# 값: 왼쪽 중요 = 1~9, 오른쪽 중요 = 1/3, 1/5, 1/7, 1/9

# 응답자1: 자립성을 가장 중요시
resp1_info = ["2026-05-18 10:00", "대학 및 출연연", "단국대학교", "연구자/교수", "지역혁신정책;연구개발", "15년 이상", "남성", "010-1234-5678"]
resp1_pairs = [
    # L1: 15쌍 (자립성vs연계성, 자립성vs실용성, ... 지속가능성vs확장성)
    3, 5, 7, 3, 5,      # 자립성 vs 나머지 (자립성 중요)
    3, 5, 1, 3,          # 연계성 vs 실용성/포용성/지속/확장
    3, 1/3, 1,           # 실용성 vs 포용성/지속/확장
    1/3, 1/3,            # 포용성 vs 지속/확장
    1,                   # 지속가능성 vs 확장성
    # L2 자립성: 3쌍
    3, 5, 3,
    # L2 연계성: 3쌍
    3, 5, 3,
    # L2 실용성: 3쌍
    1, 3, 3,
    # L2 포용성: 3쌍
    3, 1, 1/3,
    # L2 지속가능성: 3쌍
    3, 5, 3,
    # L2 확장성: 3쌍
    1/3, 1, 3,
]

# 응답자2: 실용성을 가장 중요시
resp2_info = ["2026-05-18 11:00", "기업", "삼성전자", "관리자", "기업지원;기술사업화", "10~15년", "남성", "010-2345-6789"]
resp2_pairs = [
    1/3, 1/5, 3, 1/3, 1,
    1/3, 3, 1, 3,
    5, 3, 3,
    1/3, 1/3,
    1,
    # L2 자립성
    1, 3, 3,
    # L2 연계성
    1, 3, 3,
    # L2 실용성
    3, 5, 3,
    # L2 포용성
    1, 3, 3,
    # L2 지속가능성
    1, 3, 3,
    # L2 확장성
    1/3, 3, 5,
]

# 응답자3: 균형적 시각
resp3_info = ["2026-05-18 14:00", "정부(공공기관)", "충남테크노파크", "실무자", "산학협력;산업정책", "5~10년", "여성", "010-3456-7890"]
resp3_pairs = [
    1, 3, 3, 1, 3,
    3, 3, 1, 3,
    1, 1/3, 1,
    1/3, 1/3,
    1,
    # L2 자립성
    1, 1, 1,
    # L2 연계성
    1, 1, 1,
    # L2 실용성
    1, 1, 1,
    # L2 포용성
    1, 1, 1,
    # L2 지속가능성
    1, 1, 1,
    # L2 확장성
    1, 1, 1,
]

INFO_COL_COUNT = 9  # 응답자 + 일반사항 8개
PAIR_START_COL = INFO_COL_COUNT + 1  # 10

respondents = [
    (resp1_info, resp1_pairs),
    (resp2_info, resp2_pairs),
    (resp3_info, resp3_pairs),
]

for r_idx, (info, pairs) in enumerate(respondents):
    row = 5 + r_idx  # 5, 6, 7행

    # 일반사항 (B~I열, column 2~9)
    for c_idx, val in enumerate(info):
        ws.cell(row=row, column=2 + c_idx, value=val)

    # 쌍대비교값 (J열부터)
    for c_idx, val in enumerate(pairs):
        ws.cell(row=row, column=PAIR_START_COL + c_idx, value=round(val, 4))

# 저장
output = '충남_AHP_분석_결과_v3.xlsx'
wb.save(output)
print(f"테스트 파일 생성: {output}")

# ============================================================
# 검증: Python에서 직접 AHP 계산하여 엑셀 수식 결과와 비교
# ============================================================
import numpy as np

print("\n" + "="*60)
print("Python 직접 계산으로 검증")
print("="*60)

LEVEL1 = ["자립성", "연계성", "실용성", "포용성", "지속가능성", "확장성"]
LEVEL2 = {
    "자립성": ["R&D역량", "기술흡수역량", "인재양성"],
    "연계성": ["산학연관협력", "기업네트워크", "중간지원기관"],
    "실용성": ["기술사업화", "기술애로", "매출고용"],
    "포용성": ["중소기업접근성", "청년창업", "지역사회참여"],
    "지속가능성": ["정책연속성", "거버넌스", "혁신인프라"],
    "확장성": ["외부지식기술", "글로벌시장", "신산업전환"],
}
RI_TABLE = {3: 0.58, 6: 1.24}

def calc_ahp(matrix, names, label):
    n = len(matrix)
    # Column sum
    col_sum = matrix.sum(axis=0)
    # Normalize
    norm = matrix / col_sum
    # Weights
    weights = norm.mean(axis=1)
    # Aw
    aw = matrix @ weights
    # Aw/Wi
    aw_wi = aw / weights
    # lambda_max
    lmax = aw_wi.mean()
    # CI
    ci = (lmax - n) / (n - 1) if n > 1 else 0
    # CR
    ri = RI_TABLE.get(n, 0)
    cr = ci / ri if ri > 0 else 0

    print(f"\n--- {label} ---")
    print(f"가중치: {dict(zip(names, [f'{w:.4f}' for w in weights]))}")
    print(f"λmax={lmax:.4f}, CI={ci:.4f}, RI={ri}, CR={cr:.4f}")
    print(f"판정: {'PASS' if cr <= 0.1 else 'FAIL'} (CR {'<=' if cr <= 0.1 else '>'} 0.1)")

    return weights, cr

# 기하평균 계산
all_pairs = [resp1_pairs, resp2_pairs, resp3_pairs]
geomeans = []
for i in range(len(resp1_pairs)):
    vals = [r[i] for r in all_pairs]
    gm = math.exp(sum(math.log(v) for v in vals) / len(vals))
    geomeans.append(gm)

# L1 행렬 구성 (6x6)
n1 = 6
mat1 = np.ones((n1, n1))
idx = 0
for i in range(n1):
    for j in range(i+1, n1):
        mat1[i][j] = geomeans[idx]
        mat1[j][i] = 1.0 / geomeans[idx]
        idx += 1

l1_weights, l1_cr = calc_ahp(mat1, LEVEL1, "제1계층: 핵심가치")

# L2 행렬 구성 (각 3x3)
global_weights = {}
for key in LEVEL1:
    n2 = 3
    mat2 = np.ones((n2, n2))
    for i in range(n2):
        for j in range(i+1, n2):
            mat2[i][j] = geomeans[idx]
            mat2[j][i] = 1.0 / geomeans[idx]
            idx += 1

    l2_weights, l2_cr = calc_ahp(mat2, LEVEL2[key], f"제2계층: {key}")

    ki = LEVEL1.index(key)
    for fi, factor in enumerate(LEVEL2[key]):
        gw = l1_weights[ki] * l2_weights[fi]
        global_weights[f"{key}-{factor}"] = gw

# 종합 순위
print("\n" + "="*60)
print("종합 글로벌 가중치 순위")
print("="*60)
sorted_gw = sorted(global_weights.items(), key=lambda x: -x[1])
for rank, (name, gw) in enumerate(sorted_gw, 1):
    print(f"  {rank:2d}위: {name:30s} {gw:.4f}")

print("\n엑셀 파일을 열어 위 결과와 비교하세요.")
